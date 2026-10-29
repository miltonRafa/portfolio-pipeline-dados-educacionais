from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from compatibilidade_raw import compare_content, config, Incompatible
from atualizacao_raw import ManifestUpdates
import atualizacao_raw as updates
import baixar_raw as download
import gerar_manifesto_raw as manifest


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook(path, names, historic=1.2, extra_year=None):
    book = Workbook()
    book.remove(book.active)
    for name in names:
        sheet = book.create_sheet(name)
        sheet.append(["Identificador", "Rede", *[f"VL_OBSERVADO_{y}" for y in range(2007, 2024, 2)], "VL_OBSERVADO_2025"])
        sheet.append(["Minas Gerais", "Pública (4)", *[historic] * 9, 9.9])
        if extra_year:
            sheet.cell(1, 13, f"VL_OBSERVADO_{extra_year}")
            sheet.cell(2, 13, 8.5)
    book.save(path)
    book.close()


class ContentTest(unittest.TestCase):
    def test_ideb_accepts_new_year_but_rejects_changed_history(self):
        with tempfile.TemporaryDirectory() as folder:
            a, b = Path(folder) / "a.xlsx", Path(folder) / "b.xlsx"
            sheets = [v["aba"] for v in config("src/bronze/ideb/ingest_ideb.py", "CONFIG_ABAS").values()]
            row = {"nome_local": "fonte.xlsx", "indicador": "IDEB"}
            workbook(a, sheets)
            workbook(b, sheets, extra_year=2027)
            self.assertEqual(compare_content(a, b, row)["politica"], "ideb_ate_2023")
            workbook(b, sheets, historic=1.3, extra_year=2027)
            with self.assertRaises(Incompatible) as caught:
                compare_content(a, b, row)
            self.assertEqual(caught.exception.details["codigo"], "conteudo_alterado")

    def test_support_workbooks_accept_metadata_but_reject_changed_cell_or_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            a, b = Path(folder) / "a.xlsx", Path(folder) / "b.xlsx"
            book = Workbook()
            book.active.append(["variavel", "descricao", 1])
            book.save(a)
            book.properties.title = "Metadados novos"
            book.save(b)
            row = {"nome_local": "apoio.xlsx", "indicador": "PND"}
            compare_content(a, b, row)
            book.active.cell(1, 2, "outra descricao")
            book.save(b)
            with self.assertRaises(Incompatible):
                compare_content(a, b, row)
            book.create_sheet("nova")
            book.save(b)
            with self.assertRaises(Incompatible) as caught:
                compare_content(a, b, row)
            self.assertEqual(caught.exception.details["codigo"], "abas_alteradas")
            book.close()

    def test_text_compares_every_record_and_field_without_sampling(self):
        with tempfile.TemporaryDirectory() as folder:
            a, b = Path(folder) / "a.txt", Path(folder) / "b.txt"
            a.write_bytes(b'ano;valor\r\n2025;"10"\r\n2025;20\r\n')
            b.write_bytes(b'ano;valor\n2025;10\n2025;20\n')
            row = {"nome_local": "arquivo.txt", "indicador": "PND", "ano": "2025"}
            self.assertEqual(compare_content(a, b, row)["linhas_comparadas"], 3)
            b.write_bytes(b'ano;valor\n2025;10\n2025;21\n')
            with self.assertRaises(Incompatible) as caught:
                compare_content(a, b, row)
            self.assertEqual(caught.exception.details["linha"], 3)

    def test_literal_na_is_not_confused_with_an_empty_cell(self):
        with tempfile.TemporaryDirectory() as folder:
            a, b = Path(folder) / "a.xlsx", Path(folder) / "b.xlsx"
            book = Workbook()
            book.active.append(["NA", "referencia"])
            book.save(a)
            book.active.cell(1, 1).value = None
            book.save(b)
            book.close()
            with self.assertRaises(Incompatible):
                compare_content(a, b, {"nome_local": "apoio.xlsx", "indicador": "PND"})

    def test_ideb_discovers_latest_edition_without_a_fixed_year(self):
        row = {"indicador": "IDEB", "ano": "2007-2023", "url_oficial": "https://pagina"}
        tabs = type("Tabs", (), {"tabs": ["https://pagina/2005-2025", "https://pagina/2005-2027"]})()
        latest = "https://download.inep.gov.br/ideb/resultados/divulgacao_regioes_ufs_ideb_2027.zip"
        with patch.object(download, "page", return_value=tabs), patch.object(download, "downloads", return_value=[latest]):
            self.assertEqual(download.discover(row), (latest, "https://pagina/2005-2027"))


class TransactionTest(unittest.TestCase):
    def prepare(self, folder):
        root = Path(folder)
        (root / "docs").mkdir()
        destination = root / "data/raw"
        (destination / ".download").mkdir(parents=True)
        target = destination / "pnd/fonte.txt"
        target.parent.mkdir()
        target.write_bytes(b"ano;valor\r\n2025;10\r\n")
        candidate = root / "candidato.txt"
        candidate.write_bytes(b"ano;valor\n2025;10\n")
        row = {name: "" for name in manifest.FIELDS}
        row.update(indicador="PND", ano="2025", nome_local="fonte.txt", caminho_local="data/raw/pnd/fonte.txt", sha256=digest(target))
        with (root / "docs/fontes_dados.csv").open("w", encoding="utf8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=manifest.FIELDS)
            writer.writeheader()
            writer.writerow(row)
        (root / "docs/fontes_dados.md").write_text("antes", encoding="utf8")
        (root / "docs/raw_checksums_sha256.txt").write_text("antes", encoding="utf8")
        return root, destination, target, candidate, row

    def writer(self, rows):
        (manifest.DOCS_DIR / "fontes_dados.md").write_text("depois", encoding="utf8")

    def test_equivalent_update_installs_file_hash_checksums_and_keeps_backup(self):
        with tempfile.TemporaryDirectory() as folder:
            root, destination, target, candidate, row = self.prepare(folder)
            old_hash = row["sha256"]
            with patch.object(manifest, "write_markdown", side_effect=self.writer):
                self.assertTrue(ManifestUpdates(root).accept_changed(candidate, row, destination))
            self.assertEqual(digest(target), digest(candidate))
            with (root / "docs/fontes_dados.csv").open(encoding="utf8") as file:
                self.assertEqual(next(csv.DictReader(file))["sha256"], digest(candidate))
            self.assertIn(digest(candidate), (root / "docs/raw_checksums_sha256.txt").read_text())
            self.assertEqual(row["_resultado"]["status"], "historico_igual_hash_atualizado")
            history = json.loads((destination / ".download/historico_atualizacoes.json").read_text(encoding="utf8"))
            self.assertEqual(digest(Path(history[0]["backup"])), old_hash)

    def test_download_archive_invokes_comparison_and_manifest_update(self):
        with tempfile.TemporaryDirectory() as folder:
            root, destination, target, candidate, row = self.prepare(folder)
            archive = root / "fonte.zip"
            with zipfile.ZipFile(archive, "w") as z:
                z.writestr("DADOS/fonte.txt", candidate.read_bytes())
            with patch.object(manifest, "write_markdown", side_effect=self.writer), \
                    patch.object(download, "ManifestUpdates", return_value=ManifestUpdates(root)), \
                    patch.object(download, "RemoteZip", side_effect=lambda _: archive.open("rb")):
                result = download.process_group(("https://download.inep.gov.br/fonte.zip", [row]), destination, destination / ".download", False)
            self.assertEqual(result[0]["status"], "historico_igual_hash_atualizado")
            self.assertEqual(digest(target), digest(candidate))

    def test_strict_fallback_explains_rejection_without_changing_raw(self):
        with tempfile.TemporaryDirectory() as folder:
            root, destination, target, candidate, row = self.prepare(folder)
            before = target.read_bytes()
            row["_estrito"] = True
            self.assertFalse(download.accept(candidate, row, destination))
            self.assertEqual(row["_rejeicao"]["codigo"], "hash_diferente_modo_estrito")
            self.assertIn("preservados", row["_rejeicao"]["fallback"])
            self.assertEqual(target.read_bytes(), before)

    def test_install_failure_rolls_back_raw_and_all_documents(self):
        with tempfile.TemporaryDirectory() as folder:
            root, destination, target, candidate, row = self.prepare(folder)
            before = {p: p.read_bytes() for p in (root / "docs").iterdir()}
            original = target.read_bytes()
            write = updates.atomic_bytes
            count = 0

            def fail_once(path, data):
                nonlocal count
                count += 1
                if count == 2:
                    raise OSError("Falha simulada ao gravar documentação")
                return write(path, data)

            with patch.object(manifest, "write_markdown", side_effect=self.writer), patch.object(updates, "atomic_bytes", side_effect=fail_once):
                with self.assertRaises(Incompatible) as caught:
                    ManifestUpdates(root).accept_changed(candidate, row, destination)
            self.assertEqual(caught.exception.details["codigo"], "instalacao_falhou")
            self.assertEqual(target.read_bytes(), original)
            for p, content in before.items():
                self.assertEqual(p.read_bytes(), content)
            self.assertFalse((destination / ".download/historico_atualizacoes.json").exists())

    def test_ideb_update_regenerates_real_documentation_and_source_metadata(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            docs = root / "docs"
            docs.mkdir()
            destination = root / "data/raw"
            (destination / ".download").mkdir(parents=True)
            with (ROOT / "docs/fontes_dados.csv").open(encoding="utf8", newline="") as file:
                rows = list(csv.DictReader(file))
            row = next(r for r in rows if r["indicador"] == "IDEB")
            target = root / row["caminho_local"]
            target.parent.mkdir(parents=True)
            sheets = [v["aba"] for v in config("src/bronze/ideb/ingest_ideb.py", "CONFIG_ABAS").values()]
            workbook(target, sheets)
            candidate = root / "novo.xlsx"
            workbook(candidate, sheets, extra_year=2027)
            row["sha256"] = digest(target)
            with (docs / "fontes_dados.csv").open("w", encoding="utf8", newline="") as file:
                writer = csv.DictWriter(file, fieldnames=manifest.FIELDS)
                writer.writeheader()
                writer.writerows(rows)
            (docs / "fontes_dados.md").write_text("anterior", encoding="utf8")
            (docs / "raw_checksums_sha256.txt").write_text("anterior", encoding="utf8")
            row["_url_download"] = "https://download.inep.gov.br/ideb/resultados/divulgacao_regioes_ufs_ideb_2027.zip"
            row["_membro"] = "divulgacao_regioes_ufs_ideb_2027.xlsx"
            ManifestUpdates(root).accept_changed(candidate, row, destination)
            self.assertIn("divulgacao_regioes_ufs_ideb_2027.zip", (docs / "fontes_dados.md").read_text(encoding="utf8"))
            with (docs / "fontes_dados.csv").open(encoding="utf8") as file:
                current = next(r for r in csv.DictReader(file) if r["indicador"] == "IDEB")
            self.assertIn("2027", current["anos_disponiveis"])
            self.assertIn("2007-2023", current["edicao_fonte"])

    def test_changed_history_and_missing_reference_never_update_manifest(self):
        with tempfile.TemporaryDirectory() as folder:
            root, destination, target, candidate, row = self.prepare(folder)
            before = (root / "docs/fontes_dados.csv").read_bytes()
            candidate.write_text("ano;valor\n2025;99\n", encoding="utf8")
            with self.assertRaises(Incompatible) as caught:
                ManifestUpdates(root).accept_changed(candidate, row, destination)
            self.assertEqual(caught.exception.details["codigo"], "conteudo_alterado")
            self.assertEqual((root / "docs/fontes_dados.csv").read_bytes(), before)
            target.unlink()
            with self.assertRaises(Incompatible) as caught:
                ManifestUpdates(root).accept_changed(candidate, row, destination)
            self.assertEqual(caught.exception.details["codigo"], "referencia_indisponivel")


if __name__ == "__main__":
    unittest.main()
